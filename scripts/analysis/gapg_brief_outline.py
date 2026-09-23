"""Outline for the brief to the Secretary, GA&PG. Plain Word, bullets only.

This is the **outline**, not the brief — the guide the note gets written from.
So every line is a fragment, never a sentence, and the body carries the claim
while the annexes carry the evidence for it.

    ./.venv/bin/python scripts/analysis/gapg_brief_outline.py

Three constraints drive every choice here, and each has bitten a previous
document on this project:

1. **Plain Word.** Word's own defaults, nothing else. No DPIC maroon, no font
   override, no shaded table headers, no grey italic captions. Headings are bold
   paragraphs rather than `add_heading`, because the built-in Heading styles are
   themed blue Calibri Light — which is formatting, just somebody else's. This
   file deliberately does **not** import the helpers in
   `bottleneck_report_docx.py`: those exist to impose the house style.
2. **Body of two pages or fewer.** Anything that does not fit moves to an annex.
   It never gets compressed into denser language. The first build ran to three
   pages, which is why the bullets below are fragments rather than sentences.
3. **Reader is a non-specialist.** No metric names, no statistics vocabulary, no
   tool names in the body. "One in seven filings", not "13.9%".

Output goes to `outputs/reports/`, which is gitignored: it carries figures and
counts derived from citizen grievances and is not a repository artifact.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
OUTDIR = REPO / "outputs" / "reports"

# Exhibits. `pr_27_time_by_entry.png` is the obvious choice for §2.1 and is NOT
# used: it is nearly square (4.71in tall at page width) and would take half the
# body on its own. `pr_32` is 1.69in tall, and shows *where* the time goes
# rather than only that it differs. pr_27 moves to Annex B.
EXHIBITS = {
    "pipeline": REPO / "docs/presentations/2026-08-17-value-add/assets/fallback/replay.png",
    "timing": REPO / "outputs/figures/pr_32_phases_by_entry.png",
    "hotspot": REPO / "docs/presentations/2026-08-17-value-add/assets/fallback/hotspot.png",
}
EXHIBIT_WIDTH = 3.6

# Every figure quoted in the body, with where it came from. Held here rather
# than queried so the outline builds in seconds; the same pattern
# `summary_docx.py` uses for its regression block.
#
#   rural housing, pensions  -> outputs/reports/grievance_bottlenecks_summary.docx
#   corpus scale             -> docs/PERFORMANCE.md, docs/ROADMAP.md
#   hotspot map              -> scripts/build_deck_map.py
#   outcome counts           -> janasunani/analytics/sql/closure.sql
FIGURES = {
    "corpus": "1.37 million",
    "housing_cases": "125,000",
    "housing_collector": 27,
    "housing_dept": 140,
    "reach_block_dept": 46,
    "reach_block_collector": 1,
    "pension_offices": 128,
    "pension_fastest": 1,
    "pension_slowest": 54,
    "map_filings": "1.13 million",
}


def bold_runs(para, text: str) -> None:
    """Write `text` into `para`, emboldening every "**...**" span.

    Every span, not just a leading one: a previous document on this project
    handled only the leading case and printed literal asterisks mid-sentence.
    """
    for i, part in enumerate(text.split("**")):
        if part:
            para.add_run(part).bold = bool(i % 2)


def heading(doc: Document, text: str, size: int = 11) -> None:
    """A heading as plain bold text.

    Not `add_heading`: the built-in Heading styles are themed blue Calibri
    Light, which is formatting of somebody else's choosing. Bold leaves the
    document genuinely unstyled and lets the writer apply their own.
    """
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.bold = True
    if size != 11:
        run.font.size = Pt(size)


def bullets(doc: Document, items: list, level: int = 1) -> None:
    """Terse bullets. A string is one bullet; a list is a nested level."""
    for item in items:
        if isinstance(item, list):
            bullets(doc, item, level + 1)
            continue
        style = "List Bullet" if level == 1 else f"List Bullet {min(level, 3)}"
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(0)
        bold_runs(p, item)


def note(doc: Document, text: str) -> None:
    """A bracketed instruction to the writer, not brief content."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def exhibit(doc: Document, key: str, caption: str) -> None:
    """An exhibit, at the point it belongs, with a plain caption.

    A missing file is a hard error rather than a silently absent exhibit: it
    means the notebook or deck has not been rebuilt, and a brief that quietly
    drops a figure is worse than one that refuses to build.
    """
    path = EXHIBITS[key]
    if not path.exists():
        raise SystemExit(f"missing exhibit: {path}")
    doc.add_picture(str(path), width=Inches(EXHIBIT_WIDTH))
    doc.paragraphs[-1].paragraph_format.space_before = Pt(4)
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(caption)
    run.font.size = Pt(8)


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.85)

    heading(doc, "Brief to Secretary, GA&PG — outline", size=14)
    note(doc, "Outline only. Bullets throughout. Body within two pages; detail "
              "in the annexes.")

    bullets(doc, [
        "**Architecture Plan for Janasunani 2.0 (Aug 2025) adopted into the "
        "RFP.** Five components: grievance entry · categorization and routing · "
        "monitoring and resolution · citizen engagement · analytics engine.",
        "This note: what we have since **tested**, and what we have **built**.",
    ])

    # ---------------------------------------------------------------- Part 1

    heading(doc, "PART 1 — APPLIED AI", size=12)
    heading(doc, "1.1  Automation pipeline")
    bullets(doc, [
        "Two of the five components: **grievance entry** · **categorization "
        "and routing**.",
        "Plan proposed: read a scanned or typed complaint · translate regional "
        "language · pull out key details · summarise · suggest the category · "
        "suggest where it goes · spot duplicates and repeat filings.",
        "**We built and tested this part. It is feasible.**",
        [
            f"run over the full record — {FIGURES['corpus']} complaints",
            "fast enough to sit inside the officer's workflow",
            "also tested: automatic removal of personal details",
        ],
        "Limit: Odia is the hardest part throughout.",
        "Capabilities, timings, numbers → **Annex A**",
    ])
    exhibit(doc, "pipeline",
            "One Odia scanned petition through the pipeline, with measured "
            "timings. Outputs canned, petition synthetic — the shape and the "
            "speed, not a live run.")

    heading(doc, "1.2  Voice and assisted intake")
    note(doc, "Not part of Janasunani 2.0 as being built. Forward-looking.")
    bullets(doc, [
        "**Not in the current build. Where grievance systems go next.**",
        "Why it matters, from our own records:",
        [
            "complaints arriving eight characters long",
            "three quarters of the shortest ones carry a photographed document "
            "— a letter photographed instead of typed",
            "the citizen who most needs the system is least able to type into it",
        ],
        "Changed: Odia speech recognition and synthesis now good enough. A year "
        "ago, not.",
        "Direction: citizen speaks the complaint in Odia, photographs the "
        "document.",
        "**Proposal: a rapid pilot, to establish whether it works here.**",
        "Caveat: eases access. Does **not** shorten the wait once in.",
        "Design, measures, governance → **Annex C**",
    ])

    # ---------------------------------------------------------------- Part 2

    heading(doc, "PART 2 — ANALYTICS", size=12)
    heading(doc, "2.1  Process monitoring — where the time goes")
    note(doc, "Two paragraphs to be written from these bullets.")
    bullets(doc, [
        "Plan component 2.1.3 — leadership to **identify bottlenecks and the "
        "cause of a delay**. This is what that looks like.",
        "Today: the system counts complaints closed. Cannot say **where a "
        "complaint waited**.",
        f"**Rural housing, last year, {FIGURES['housing_cases']} cases:**",
        [
            f"to the District Collector → **{FIGURES['housing_collector']} "
            f"days** · to the department → **{FIGURES['housing_dept']} days**",
            f"reason: **{FIGURES['reach_block_dept']} days just to reach the "
            f"block office that does the work**. Via the Collector, "
            f"**{FIGURES['reach_block_collector']} day**. Nobody in the field "
            f"in that time.",
        ],
        f"**Pensions: where it is sent makes no difference.** Which block "
        f"office handles it does — across {FIGURES['pension_offices']} offices "
        f"the average wait runs **{FIGURES['pension_fastest']} day to "
        f"{FIGURES['pension_slowest']}**.",
        "Two departments, one system, two different problems. One statewide "
        "dashboard shows neither.",
        "**One in seven filings: the same complaint sent again.** Clock "
        "restarts, someone must close it.",
        "**A quarter of pension cases record no block office.** Unroutable, "
        "invisible to their district, slowest group.",
        "All from records already held. No new collection. No extra work for "
        "any officer.",
        "Caveat: measures how long a case stayed open, not effort. **No office "
        "ranking presented.**",
        "Method, full results → **Annex B**",
    ])
    exhibit(doc, "timing",
            "Average days at each step, by where the complaint was first sent. "
            "Steps mix travel between offices with field work.")

    heading(doc, "2.2  Governance x-ray — the problem behind the complaints")
    bullets(doc, [
        "Plan component 2.1.5, the **Analytics Engine** — grouping grievances "
        "by place and time, hotspot maps, an **Impact Opportunity** view of how "
        "many citizens benefit if a cluster is fixed. Part built.",
        "System closes complaints one at a time. Most are versions of a "
        "**smaller number of real problems**.",
        f"**Hotspots — built.** Group complaints that are really the same "
        f"problem; see where one is concentrated and growing. Map covers "
        f"{FIGURES['map_filings']} filings, district level (→ Annex B).",
        [
            "block level needs a boundary map, a tidy-up of block name "
            "spellings, a decision to report at that level. **No new "
            "technology. Nothing new asked of citizens.**",
        ],
        "**Outcomes — the biggest blind spot. The plan already fixes it** "
        "(2.1.3, structured closing statuses).",
        [
            "today: **fewer than one in seventeen closed cases records a "
            "benefit actually given**. So the system is manageable only on "
            "speed — and speed alone is met by closing faster.",
        ],
        "**Equity — recorded, never used.** Gender, disability, remoteness all "
        "held; never checked against waits or outcomes. Nothing is designed to "
        "discriminate — which is when unequal outcomes go unnoticed.",
        "**Repeat filing, read the other way.** A repeat = the first attempt "
        "failed. A service measure, no new data. Caveat: not returning ≠ satisfied.",
        "Still needed: block level · the structured closing status · citizen "
        "feedback (2.1.4; none exists today).",
    ])
    # ---------------------------------------------------------------- Annexes

    doc.add_page_break()
    heading(doc, "ANNEXES", size=12)
    note(doc, "All detail here. Exact counts, samples, methods, tool names.")

    heading(doc, "Annex A — pipeline capabilities and measured performance")
    bullets(doc, [
        "Our own tested pipeline — a superset of what is in the plan.",
        "Six stages:",
        [
            "page discovery and format label",
            "text extraction (OCR)",
            "PII redaction — name, mobile, landline, Aadhaar, PAN, bank "
            "account, scheme ID, email; typed tokens with character spans",
            "page-type gate — letters and applications vs ID cards and bills",
            "summarisation",
            "category classification",
        ],
        "Three capabilities beside the pipeline:",
        [
            "routing suggestion — district-aware, learned from case history",
            "duplicate and repeat-filing detection — separate salted path for "
            "the same citizen returning",
            "advisory triage — flags, never blocks",
        ],
        "One API call returns redaction, category, summary, routing, triage.",
        "Scale: 1,371,288 complaints · 6,556,171 action rows.",
        "Timings: typed complaint under a second warm · scanned document about "
        "13 seconds.",
        "Stage by stage: what was measured, on what sample, when.",
        "Gates that currently fail, and why they were not relaxed.",
    ])

    heading(doc, "Annex B — process monitoring and hotspot detail")
    bullets(doc, [
        "Sample definitions · the five specifications · route and office tables.",
        "pr_27_time_by_entry.png — time to close by entry point, as distributions.",
        "Pensions office-spread table: role · offices · average · fastest · "
        "slowest.",
        "Exact case counts behind every figure quoted in the body.",
        "Caveats in full, including that step boundaries contain travel between "
        "offices and field work together.",
        "Rule we hold to: 500 people complaining about one road is **the "
        "signal**, not duplication to remove. Spikes are labelled, never "
        "suppressed.",
        "Block detail: recorded on **82.7% of filings across 461 district-block "
        "pairs**. Needs a public block boundary file, a crosswalk over **427 "
        "block-name spellings**, and admitting block as a reporting dimension.",
        "Outcome counts: **25,182 of 449,085** closed cases (2024-25) record a "
        "benefit. **60.9%** of **776,922** template closures use wording "
        "claiming no action.",
        [
            "Qualifier: that share **rises** with work done — 64.8% at six or "
            "more steps. A recording problem, not a performance finding.",
        ],
    ])
    exhibit(doc, "hotspot",
            "Hotspot map — 1.13 million filings, one dot per district per "
            "theme, dot size by count. District level is the current limit. A "
            "dot map, not a shaded one: shading a district by volume just "
            "re-reports population.")

    heading(doc, "Annex C — voice pilot design")
    bullets(doc, [
        "Channels — one pilot, two channels, two questions:",
        [
            "phone line — tests reach; a feature phone is enough",
            "SMS link for the document photo — tests the document half",
        ],
        "Measures:",
        [
            "Odia transcription quality against hand-checked transcripts — no "
            "such set exists; the pilot creates one",
            "completion rate",
            "does a spoken complaint route and categorise as well as a typed one",
            "is the uploaded document legible enough",
            "who uses it, against the profile of existing filers",
        ],
        "Scope to fix with the department: one department · 1–2 districts · "
        "defined window · target call count.",
        "Governance, before any audio moves:",
        [
            "consent at call start — a recording identifies more than text",
            "a third-party speech vendor is an external-data decision; gated "
            "and logged",
            "retention — how long audio is kept, or whether it is kept at all "
            "after transcription",
            "a failed transcription must never become a dropped complaint",
        ],
        "Open unknowns: Odia speech recognition on rural dialect and "
        "phone-quality audio · workflow and prompt design · data privacy.",
    ])

    OUTDIR.mkdir(parents=True, exist_ok=True)
    out = OUTDIR / "gapg_brief_outline.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")
